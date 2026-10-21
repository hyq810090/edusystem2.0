"""讲义 Word export, run inside a space's worker.

The same layout as the PDF (templates/handout/mathbank-handout.sty): accent
colour, heading hierarchy, per-section numbering (each kind counts on its own), tinted definition
boxes, framed theorems, side-ruled notes. Questions go through the engine's
Word paper builder, so formulas are editable OMML and figures and choices are
laid out as in papers. Fonts are the ones teachers' Word has: 宋体 and Times
New Roman for text, 黑体 for labels and headings.
"""
import re
from io import BytesIO
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from platform_app.handout_export import FIGURE_CM, FIGURE_LIMITS, LIST_ITEM, NOTES, THEOREMS

HEX = {'mbAccent': 'A51C30', 'mbInk': '212529', 'mbMuted': '6C757D'}
COLOURS = {name: RGBColor.from_string(value) for name, value in HEX.items()}
ACCENT, INK, MUTED = COLOURS['mbAccent'], COLOURS['mbInk'], COLOURS['mbMuted']
TINT, PANEL, RULE = 'FAF4F5', 'F5F6F8', 'DADCE0'
MARGIN_CM = 2.2
TEXT_CM = 21 - 2 * MARGIN_CM
WIDTH_TWIPS = int(TEXT_CM / 2.54 * 1440)
IMAGE = re.compile(r'!\[.*?\]\(([^)]+)\)')
# w:pPr children that must follow w:pBdr (ECMA-376 order); Word rejects out-of-order properties.
AFTER_PBDR = ('w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
              'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
              'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
              'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId',
              'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')


def lighten(value, strength):
    """The PDF's `colour!55`: the colour mixed with white."""
    channels = (int(value[i:i + 2], 16) for i in (0, 2, 4))
    return ''.join(f'{round(c * strength + 255 * (1 - strength)):02X}' for c in channels)


def _edge(parent, edge, size, color, space):
    node = OxmlElement(f'w:{edge}')
    if size:
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), str(size))
        node.set(qn('w:space'), str(space))
        node.set(qn('w:color'), color)
    else:
        node.set(qn('w:val'), 'nil')
    parent.append(node)


def paragraph_border(paragraph, edge, size, color):
    """One rule above or below a paragraph (size in eighths of a point)."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    _edge(borders, edge, size, color, 4)
    p_pr.insert_element_before(borders, *AFTER_PBDR)


def cell_style(cell, fill=None, **edges):
    """Borders and shading of a one-cell box; call before the cell margins (tcPr order)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        size, color = edges.get(edge, (0, ''))
        _edge(borders, edge, size, color, 0)
    tc_pr.append(borders)
    if fill:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:color'), 'auto')
        shading.set(qn('w:fill'), fill)
        tc_pr.append(shading)


def build(engine, payload):
    from mathbank import word_export_helper as w
    from mathbank.paper_helper import collect_referenced_images

    teacher = payload['version'] == 'teacher'
    document = payload['document']
    subject_line = f"{payload['stage']}{payload['subject']}"

    class Handout(w.WordExamBuilder):
        lead = None  # an example's label, written where papers write the question number

        def _configure_document(self):
            super()._configure_document()
            section = self.doc.sections[0]
            section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.2)
            section.left_margin = section.right_margin = Cm(MARGIN_CM)
            section.different_first_page_header_footer = True
            header = section.header.paragraphs[0]
            tabs = header.paragraph_format.tab_stops
            # The Header style's own centre and right stops (3.25 in, 6.5 in) would catch the tab first.
            tabs.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CLEAR)
            tabs.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.CLEAR)
            tabs.add_tab_stop(Cm(TEXT_CM), WD_TAB_ALIGNMENT.RIGHT)
            w._set_run_font(header.add_run(f"{payload['title']}\t{subject_line} · 讲义"), 9,
                            color=MUTED, cjk_font=w.CJK_HEADING_FONT)
            paragraph_border(header, 'bottom', 4, RULE)
            for footer in (section.footer.paragraphs[0], section.first_page_footer.paragraphs[0]):
                for run in list(footer._p.findall(qn('w:r'))):
                    footer._p.remove(run)
                footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._append_field(footer, 'PAGE')

        def add_question(self, item, number):
            self.current = item.question
            super().add_question(item, number)

        def _add_right_figure_row(self, path, *, max_width, max_height):
            # The engine's stem/figure row with the handout's own figure limits (as its PDF minipage).
            width, height = FIGURE_LIMITS.get(self.current.get('figure_size'), FIGURE_LIMITS['auto'])
            return super()._add_right_figure_row(path, max_width=width / 2.54, max_height=height / 2.54)

        def _format_question_paragraph(self, paragraph, number):
            if self.lead is None:
                return super()._format_question_paragraph(paragraph, number)
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(3)
            label(paragraph, self.lead, ACCENT)

    def label(paragraph, text, color, size=w.BODY_FONT_SIZE):
        w._set_run_font(paragraph.add_run(text + '  '), size, bold=True, color=color, cjk_font=w.CJK_HEADING_FONT)

    builder = Handout({}, w.WordExportDiagnostics())
    doc = builder.doc
    uploads = Path(engine.UPLOAD_DIR)

    # Every text that may carry formulas, converted to OMML in one Pandoc pass;
    # every image the teacher's own text refers to, resolved inside the uploads folder.
    texts = [payload['title'], document['kicker'], document['subtitle'], *document['objectives']]
    items, sources = {}, []
    for section in document['sections']:
        texts.append(section['title'])
        for block in section['blocks']:
            texts += [block.get(key) or '' for key in ('title', 'body', 'caption', 'solution')]
            if block['kind'] == 'figure' and block['image']:
                sources.append({'question': {'content': f"![]({block['image']})"}})
            for question in ([block['question']] if block['kind'] == 'example' else block.get('questions', [])):
                item = items[id(question)] = w._prepare_question({'question': question, 'score': 0}, uploads)
                texts += [item.stem, item.answer, *item.choices]
    texts = [text for text in texts if text]
    sources += [{'question': {'content': text}} for text in texts if IMAGE.search(text)]
    images = [Path(path) for path in collect_referenced_images(sources, engine.UPLOAD_DIR, engine.UPLOAD_DIR_REL)]
    builder.omml_map = w.PandocOmmlConverter().convert_many(w._extract_formulas(texts))

    def rich(paragraph, text, parent=None):
        """Teacher prose: **bold** runs, maths and images; tables go through the paper builder."""
        if '**' not in text:
            builder._add_content_blocks(paragraph, text, parent=parent, images=images)
            return
        for index, part in enumerate(text.split('**')):
            before = len(paragraph.runs)
            builder.add_mixed(paragraph, part, images=images)
            if index % 2:
                for run in paragraph.runs[before:]:
                    run.bold = True

    def prose(text, first=None, parent=None, indent=True):
        """Blank-line separated paragraphs and '- ' / '1. ' list lines; `first` continues a label's paragraph."""
        container = parent or doc
        state = {'first': first}

        def paragraph_with(content, marker=None):
            if not content:
                return
            paragraph = state['first'] if state['first'] is not None else container.add_paragraph()
            state['first'] = None
            paragraph.paragraph_format.space_after = Pt(3)
            if marker:
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.first_line_indent = Cm(-0.45)
                w._set_run_font(paragraph.add_run(marker + ' '), w.BODY_FONT_SIZE, color=MUTED)
            elif indent and not paragraph.runs:
                paragraph.paragraph_format.first_line_indent = Cm(0.74)
            rich(paragraph, content, parent)

        for chunk in re.split(r'\n\s*\n', (text or '').replace('\r\n', '\n').strip()):
            lines = []
            for line in chunk.split('\n'):
                match = LIST_ITEM.match(line)
                if match:
                    paragraph_with(' '.join(lines).strip())
                    lines = []
                    paragraph_with(match.group(3).strip(), '•' if match.group(1) else match.group(2) + '.')
                else:
                    lines.append(line.strip())
            paragraph_with(' '.join(lines).strip())

    def box(fill=None, **edges):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        w._set_table_geometry(table, [WIDTH_TWIPS])
        cell = table.cell(0, 0)
        cell_style(cell, fill, **edges)
        w._set_cell_margins(cell, 110, 170, 110, 170)
        gap(5)  # also keeps two boxes in a row from merging into one table
        return cell

    def gap(points):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        spacer.paragraph_format.line_spacing = Pt(points)
        return spacer

    def heading(text, size, color=INK, before=14, after=6):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.keep_with_next = True
        if text:
            w._set_run_font(paragraph.add_run(text), size, bold=True, color=color, cjk_font=w.CJK_HEADING_FONT)
        return paragraph

    # Heading: lesson label, title, subtitle, accent rule, subject line.
    if document['kicker']:
        heading(document['kicker'], 11, ACCENT, before=0, after=3)
    heading(payload['title'], 22, before=0, after=4)
    if document['subtitle']:
        w._set_run_font(doc.add_paragraph().add_run(document['subtitle']), 12.5, color=MUTED)
    meta = doc.add_paragraph()
    paragraph_border(meta, 'top', 7, HEX['mbAccent'])
    meta.paragraph_format.space_before = Pt(8)
    meta.paragraph_format.space_after = Pt(12)
    textbook = f" · {payload['textbook']}" if payload.get('textbook') else ''
    w._set_run_font(meta.add_run(f"{subject_line}{textbook} · {'教师版' if teacher else '学生版'}"), 9,
                    color=MUTED, cjk_font=w.CJK_HEADING_FONT)

    objectives = [item for item in document['objectives'] if item.strip()]
    if objectives:
        cell = box(PANEL)
        label(cell.paragraphs[0], '学习目标', ACCENT)
        for index, item in enumerate(objectives, 1):
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.first_line_indent = Cm(-0.45)
            w._set_run_font(paragraph.add_run(f'{index}  '), 9, color=MUTED)
            rich(paragraph, item)

    answers = []
    for number, section in enumerate(document['sections'], 1):
        paragraph = heading('', 15, before=18, after=6)
        w._set_run_font(paragraph.add_run(f'{number}  '), 15, bold=True, color=ACCENT, cjk_font=w.CJK_HEADING_FONT)
        w._set_run_font(paragraph.add_run(section['title']), 15, bold=True, color=INK, cjk_font=w.CJK_HEADING_FONT)
        counters = {}  # each kind counts on its own: 定义 1.1, 1.2 …; 例 1.1 …
        for block in section['blocks']:
            kind = block['kind']
            if kind in {'definition', *THEOREMS, 'example', 'figure'}:
                counters[kind] = counters.get(kind, 0) + 1
            tag = f'{number}.{counters.get(kind, 0)}'
            titled = f"{tag}  {block['title']}" if block.get('title') else tag
            if kind == 'text':
                prose(block['body'])
            elif kind == 'definition':
                cell = box(TINT, left=(18, HEX['mbAccent']))
                label(cell.paragraphs[0], f'定义 {titled}', ACCENT)
                prose(block['body'], cell.paragraphs[0], cell, indent=False)
            elif kind in THEOREMS:
                cell = box(None, **{edge: (5, RULE) for edge in ('top', 'left', 'bottom', 'right')})
                label(cell.paragraphs[0], f'{THEOREMS[kind]} {titled}', INK)
                prose(block['body'], cell.paragraphs[0], cell, indent=False)
            elif kind == 'proof':
                first = doc.add_paragraph()
                label(first, '证明', MUTED)
                prose(block['body'] + '  □', first, indent=False)
            elif kind in NOTES:
                name, colour = NOTES[kind]
                cell = box(None, left=(8, lighten(HEX[colour], 0.55)))
                label(cell.paragraphs[0], name, COLOURS[colour], 9)
                prose(block['body'], cell.paragraphs[0], cell, indent=False)
            elif kind == 'summary':
                cell = box(PANEL, top=(6, HEX['mbAccent']), bottom=(6, HEX['mbAccent']))
                label(cell.paragraphs[0], '本节小结', ACCENT)
                prose(block['body'], parent=cell, indent=False)
            elif kind == 'figure':
                path = next((image for image in images if image.name == Path(block['image']).name), None)
                if path:
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(6)
                    paragraph.paragraph_format.keep_with_next = True
                    width, height = FIGURE_CM[block['size']]
                    builder._add_image(paragraph, path, max_width=width / 2.54, max_height=height / 2.54)
                    caption = doc.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.paragraph_format.space_after = Pt(6)
                    text = f'图 {tag}' + (f"  {block['caption']}" if block['caption'] else '')
                    w._set_run_font(caption.add_run(text), 9, color=MUTED, cjk_font=w.CJK_HEADING_FONT)
            elif kind == 'example':
                item = items[id(block['question'])]
                builder.lead = f'例 {titled}'
                builder.add_question(item, 0)
                builder.lead = None
                if teacher or document['settings']['student_solutions'] == 'show':
                    cell = box(None, left=(6, RULE))
                    first = cell.paragraphs[0]
                    label(first, '解', ACCENT)
                    if (block.get('solution') or '').strip():
                        prose(block['solution'], first, cell, indent=False)
                    elif item.answer.strip():
                        builder._add_content_blocks(first, item.answer, parent=cell, images=images)
                    else:
                        w._set_run_font(first.add_run('（暂无答案）'), w.BODY_FONT_SIZE, color=MUTED)
                elif block['space']:
                    gap(1).paragraph_format.space_after = Cm(block['space'])
            elif kind == 'exercises' and block['questions']:
                title = block['title'] or '练习'
                heading(title, 12, before=12, after=4)
                group = [items[id(question)] for question in block['questions']]
                for index, item in enumerate(group, 1):
                    builder.add_question(item, index)
                if teacher:
                    answers.append((f'{number}  {title}', group))

    if answers:
        paragraph = heading('参考答案', 15, before=24, after=6)
        paragraph_border(paragraph, 'top', 7, HEX['mbAccent'])
        for title, group in answers:
            heading(title, w.BODY_FONT_SIZE, ACCENT, before=8, after=2)
            for index, item in enumerate(group, 1):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(3)
                w._set_run_font(paragraph.add_run(f'{index}. '), w.BODY_FONT_SIZE, bold=True)
                if item.answer.strip():
                    builder._add_content_blocks(paragraph, item.answer, images=images)
                else:
                    w._set_run_font(paragraph.add_run('（暂无答案）'), w.BODY_FONT_SIZE, color=MUTED)

    properties = doc.core_properties
    properties.title = payload['title']
    properties.subject = f'{subject_line}讲义'
    properties.author = 'MathBank'
    output = BytesIO()
    doc.save(output)
    return output.getvalue(), None
