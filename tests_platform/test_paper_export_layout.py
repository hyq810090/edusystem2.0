"""Paper exports: a 题干右侧 figure in Word sits beside the stem with the choices below (as the PDF's
minipages), and the stem's own empty brackets go even before a full stop, so \\paren prints one pair."""
from io import BytesIO
from pathlib import Path
import json
import shutil
import subprocess
import sys
import zipfile

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'engine'))
from mathbank.paper_helper import clean_choice_stem_parentheses, clean_content_for_latex  # noqa: E402
from mathbank.word_export_helper import build_word_document  # noqa: E402

CHOICES = r'\begin{choices}\item $1$\item $2$\item $3$\item $4$\end{choices}'
FIGURE = '![](/static/uploads/right.png)'
BRACKETS = [('则阴影部分的面积为（ ）。', '则阴影部分的面积为'), ('值为 $(\\quad)$．', '值为'),
            ('值为（　）.', '值为'), ('下列正确的是( )', '下列正确的是')]
KEPT = ['求值（结果保留π）。', '已知 x=1。', '共有（3）种。']


def paper(tmp_path, content, question_type='single_choice'):
    Image.new('RGB', (300, 240), 'white').save(tmp_path / 'right.png')
    blob, _ = build_word_document('右图', '', 'exam', [{'question': {
        'id': 1, 'content': content, 'question_type': question_type, 'answer_markdown': '',
        'figure_align': 'right', 'image_paths': ['/static/uploads/right.png']}, 'score': 5}], uploads_dir=tmp_path)
    return blob, Document(BytesIO(blob))


def test_word_right_figure_sits_beside_the_stem_with_full_width_choices(tmp_path):
    blob, document = paper(tmp_path, f'如图，求阴影部分的面积（ ）。\n\n{FIGURE}\n{CHOICES}')
    # A floating picture made Microsoft Word push the choices under it and overlap the next figure.
    assert '<wp:anchor' not in zipfile.ZipFile(BytesIO(blob)).read('word/document.xml').decode()
    row, choices = document.tables
    left, right = row.rows[0].cells
    assert left.text.startswith('1. ') and '阴影部分的面积' in left.text and not left._tc.xpath('.//w:drawing')
    extents = right._tc.xpath('.//wp:extent')
    assert len(extents) == 1 and int(extents[0].get('cx')) == pytest.approx(Inches(1.55), rel=0.01)
    stem_width, figure_width = (int(width) for width in row._tbl.xpath('./w:tblGrid/w:gridCol/@w:w'))
    assert 8900 <= stem_width + figure_width <= 9100 and figure_width < stem_width
    assert choices.cell(0, 0).text.strip().startswith('A.') and len(choices.columns) == 4
    # A paragraph between the row and the choices keeps Word from joining the two tables.
    assert row._tbl.getnext().tag.endswith('}p')


def test_word_stem_with_a_table_keeps_full_width_and_puts_the_figure_below(tmp_path):
    table = '\\begin{tabular}{|c|c|}\n\\hline 甲 & 乙 \\\\\n\\hline\n\\end{tabular}'
    blob, document = paper(tmp_path, f'数据如下。\n\n{table}\n\n求两数之和。\n\n{FIGURE}', 'detailed_answer')
    assert len(document.tables) == 1 and not document.tables[0]._tbl.xpath('.//w:drawing')
    xml = zipfile.ZipFile(BytesIO(blob)).read('word/document.xml').decode()
    assert '<wp:anchor' not in xml and xml.count('<w:drawing>') == 1
    figure = next(p for p in document.paragraphs if p._p.xpath('.//w:drawing'))
    assert figure.alignment == 2  # right, as bottom_right


def test_stem_brackets_before_a_full_stop_are_left_to_paren():
    for stem, expected in BRACKETS:
        assert clean_choice_stem_parentheses(stem) == expected, stem
    for stem in KEPT:
        assert clean_choice_stem_parentheses(stem) == stem
    tex = clean_content_for_latex('则阴影部分的面积为（ ）。\n' + CHOICES, 'single_choice')
    assert tex.count(r'\paren') == 1 and '（ ）' not in tex and '面积为' in tex


@pytest.mark.skipif(not shutil.which('node'), reason='node is not installed')
def test_preview_cleaner_follows_the_same_rule(tmp_path):
    source = (ROOT/'engine/static/js/editor.js').read_text(encoding='utf-8')
    start = source.index('function cleanChoiceStemParentheses(text)')
    marker = 'window.preprocessFormulaForKaTeX = preprocessFormulaForKaTeX;'
    cases = [stem for stem, _ in BRACKETS] + KEPT
    script = tmp_path/'clean.js'
    script.write_text('global.window = {MathBankSafe: {safeImageUrl(v) { return v; }, escapeAttribute(v) { return v; }}};\n'
                      + source[start:source.index(marker, start) + len(marker)]
                      + f'\nconsole.log(JSON.stringify({json.dumps(cases, ensure_ascii=False)}.map(cleanChoiceStemParentheses)));\n',
                      encoding='utf-8')
    result = subprocess.run(['node', str(script)], capture_output=True, encoding='utf-8', check=True)
    assert json.loads(result.stdout) == [expected for _, expected in BRACKETS] + KEPT
