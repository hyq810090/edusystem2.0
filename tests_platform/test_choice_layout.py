"""Options per row: an authored choices[columns=N], else exam-zh's fit-and-halve rule in Word."""
from io import BytesIO
from pathlib import Path
import sys

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'engine'))
from mathbank.choice_layout import authored_columns, estimate_columns  # noqa: E402
from mathbank.word_export_helper import build_word_document  # noqa: E402


def test_authored_columns_reads_the_exam_zh_key():
    assert authored_columns(r'题干\begin{choices}[columns=2]\item 甲\end{choices}') == 2
    assert authored_columns(r'\begin{choices} [label-pos=left, columns=4]\item 甲\end{choices}') == 4
    assert authored_columns(r'\begin{choices}[columns=3]\item 甲\end{choices}') == 3
    for automatic in (r'\begin{choices}\item 甲\end{choices}', r'\begin{choices}[columns=9]\item 甲\end{choices}',
                      r'\begin{choices}[label-pos=left]\item 甲\end{choices}', '没有选项'):
        assert authored_columns(automatic) is None


def test_estimate_follows_the_printed_width_of_the_widest_option():
    # Short formulas and image options: one row of four.
    assert estimate_columns([r'$\frac{2}{3}\pi$', r'$\frac{8}{3}\pi$', r'$6\pi$', r'$\frac{10}{3}\pi$']) == 4
    assert estimate_columns(['![](/static/uploads/1/a.png)'] * 4) == 4
    # A fraction prints as wide as its longer line, not as long as its source.
    assert estimate_columns([r'$\frac{2\Delta x(t_{1}-t_{2})}{t_{1}t_{2}(t_{1}+t_{2})}$'] * 4) == 4
    # Sixteen characters fit half the line; twenty do not.
    assert estimate_columns(['甲、乙在同一时刻总能到达同一高度', '甲比乙先到达终点']) == 2
    assert estimate_columns([r'在 $0$—$10\mathrm{s}$ 内两车逐渐靠近', '乙']) == 2
    assert estimate_columns(['两车刹车时相距的距离一定小于一百米才会相撞', '甲']) == 1
    assert estimate_columns([]) == 1


def word_option_columns(content):
    blob, _ = build_word_document('选项排列', '', 'exam', [{'question': {
        'id': 1, 'content': content, 'question_type': 'single_choice', 'answer_markdown': ''}, 'score': 5}])
    for table in Document(BytesIO(blob)).tables:
        if table.cell(0, 0).text.strip().startswith('A.'):
            return len(table.columns), len(table.rows)
    raise AssertionError('no choices table')


def test_word_honours_an_authored_layout_and_estimates_otherwise():
    options = r'\item 甲\item 乙\item 丙\item 丁\end{choices}'
    assert word_option_columns('下列正确的是\\begin{choices}' + options) == (4, 1)
    assert word_option_columns('下列正确的是\\begin{choices}[columns=2]' + options) == (2, 2)
    assert word_option_columns('下列正确的是\\begin{choices}[columns=1]' + options) == (1, 4)
