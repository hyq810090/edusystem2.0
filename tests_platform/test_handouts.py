"""讲义: storage, permissions, question snapshots, lesson folders, exports and the migration."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import unquote
import json
import sqlite3
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from test_platform import create_user, login_user, new_space, question, system  # noqa: F401
from platform_app import migrations, store

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'engine'))


def handouts(space):
    return f'/platform/api/spaces/{space}/handouts'


def sample_document(example, exercises):
    return {
        'kicker': '第 3 讲', 'subtitle': '概念与计算', 'objectives': ['理解 $\\sin x$ 的定义。'],
        'settings': {'student_solutions': 'blank'},
        'sections': [
            {'title': '概念', 'blocks': [
                {'kind': 'text', 'body': '引入：100% 的学生 & 老师都会用到。'},
                {'kind': 'definition', 'title': '正弦', 'body': '**正弦**是对边比斜边。\n- 第一条\n- 第二条'},
                {'kind': 'theorem', 'title': '', 'body': '$\\sin^2x+\\cos^2x=1$'},
                {'kind': 'warning', 'body': '注意符号。'},
                {'kind': 'example', 'title': '基础', 'question': example, 'solution': '由定义得 ____ 。', 'space': 3},
                {'kind': 'proof', 'body': '显然。'},
            ]},
            {'title': '练习', 'blocks': [
                {'kind': 'exercises', 'title': '基础练习', 'questions': exercises},
                {'kind': 'summary', 'body': '- 定义\n- 定理'},
            ]},
        ],
    }


def picture(space, name):
    """A real file in the space's uploads, which the bank requires before a question may show it."""
    uploads = store.DATA / 'banks' / str(space) / 'static' / 'uploads'
    uploads.mkdir(parents=True, exist_ok=True)
    Image.new('RGB', (300, 240), 'white').save(uploads / name)
    return f'/static/uploads/{name}'


def snapshots(client, space, ids):
    response = client.post(handouts(space) + '/snapshots', json={'ids': ids})
    assert response.status_code == 200, response.text
    return response.json()


def test_handout_crud_revision_and_validation(system):  # noqa: F811
    space = new_space(system, '讲义存取')['id']
    other = new_space(system, '讲义存取二')['id']
    assert system.post(handouts(space), json={'title': '  '}).status_code == 400
    created = system.post(handouts(space), json={'title': '第一讲'}).json()
    assert created['revision'] == 1 and [s['title'] for s in created['document']['sections']] == ['知识梳理']
    assert created['document']['settings'] == {'student_solutions': 'blank'}
    hid = created['id']
    listed = system.get(handouts(space)).json()
    assert [h['id'] for h in listed] == [hid] and 'document' not in listed[0]
    assert listed[0]['summary'] == {'kicker': '', 'sections': 1, 'blocks': 0}
    # Whole-document save with optimistic revision; unknown fields are dropped, ids assigned.
    document = {'kicker': '第 1 讲', 'extra': 'dropped', 'sections': [
        {'title': '小节', 'blocks': [{'kind': 'definition', 'title': '集合', 'body': '对象的全体。', 'colour': 'red'}]}]}
    saved = system.put(f'{handouts(space)}/{hid}', json={'title': '第一讲 集合', 'document': document, 'revision': 1})
    assert saved.status_code == 200, saved.text
    body = saved.json()
    assert body['revision'] == 2 and body['title'] == '第一讲 集合' and 'extra' not in body['document']
    block = body['document']['sections'][0]['blocks'][0]
    assert block['id'].startswith('b') and 'colour' not in block and block['body'] == '对象的全体。'
    stale = system.put(f'{handouts(space)}/{hid}', json={'title': '覆盖', 'document': document, 'revision': 1})
    assert stale.status_code == 409
    # Ids the client already has survive a save, so the editor keeps its selection.
    again = system.put(f'{handouts(space)}/{hid}', json={'title': '第一讲 集合', 'document': body['document'], 'revision': 2}).json()
    assert again['document']['sections'][0]['id'] == body['document']['sections'][0]['id']
    assert again['document']['sections'][0]['blocks'][0]['id'] == block['id']
    # Validation.
    for bad in ({'sections': [{'title': 'x', 'blocks': [{'kind': 'chart'}]}]},
                {'objectives': ['目标'] * 13},
                {'sections': [{'title': 'x', 'blocks': [{'kind': 'example', 'question': {'content': '无来源'}}]}]},
                {'sections': [{'title': 'x' * 121, 'blocks': []}]},
                'not a document'):
        response = system.put(f'{handouts(space)}/{hid}', json={'title': 't', 'document': bad, 'revision': 3})
        assert response.status_code == 400, bad
    # Rename, copy, other space, delete.
    assert system.patch(f'{handouts(space)}/{hid}', json={'title': '集合'}).json()['title'] == '集合'
    copy = system.post(f'{handouts(space)}/{hid}/copy').json()
    assert copy['title'] == '集合（副本）' and copy['revision'] == 1 and copy['summary']['blocks'] == 1
    assert system.get(f'{handouts(other)}/{hid}').status_code == 404
    assert system.put(f'{handouts(other)}/{hid}', json={'title': 't', 'document': {}, 'revision': 3}).status_code == 404
    assert system.delete(f'{handouts(space)}/{hid}').status_code == 200
    assert system.get(f'{handouts(space)}/{hid}').status_code == 404
    assert [h['id'] for h in system.get(handouts(space)).json()] == [copy['id']]


def test_viewers_read_and_export_but_do_not_write(system):  # noqa: F811
    space = new_space(system, '讲义权限', 'team')['id']
    hid = system.post(handouts(space), json={'title': '共享讲义'}).json()['id']
    create_user(system, 'handout_viewer')
    viewer = login_user('handout_viewer', system)
    assert [h['id'] for h in viewer.get(handouts(space)).json()] == [hid]
    assert viewer.get(f'{handouts(space)}/{hid}').status_code == 200
    assert viewer.get(f'{handouts(space)}/{hid}/export', params={'format': 'docx'}).status_code == 200
    assert viewer.post(handouts(space), json={'title': '越权'}).status_code == 403
    assert viewer.put(f'{handouts(space)}/{hid}', json={'title': 'x', 'document': {}, 'revision': 1}).status_code == 403
    assert viewer.patch(f'{handouts(space)}/{hid}', json={'title': 'x'}).status_code == 403
    assert viewer.post(f'{handouts(space)}/{hid}/copy').status_code == 403
    assert viewer.delete(f'{handouts(space)}/{hid}').status_code == 403
    assert viewer.post(handouts(space) + '/snapshots', json={'ids': [1]}).status_code == 403
    private = new_space(system, '讲义私有')['id']
    assert viewer.get(handouts(private)).status_code == 404


def test_snapshots_keep_the_question_and_report_bank_changes(system):  # noqa: F811
    space = new_space(system, '讲义快照')['id']
    image = picture(space, 'fig-a.png')
    created = question(system, space, content=f'如图 ![]({image}) 求 $x$。', answer_markdown='$x=1$',
                       image_paths=json.dumps([image]))
    assert created.status_code == 200, created.text
    first = created.json()['question']
    second = question(system, space, content='计算 $1+1$。').json()['question']
    taken = snapshots(system, space, [second['id'], first['id'], 99999])
    assert [s['source_id'] for s in taken] == [second['id'], first['id']]
    assert all(len(s['source_hash']) == 64 and 'id' not in s for s in taken)
    # Browser form outwards, bank form when stored.
    assert f'/static/uploads/{space}/fig-a.png' in taken[1]['content']
    assert system.post(handouts(space) + '/snapshots', json={'ids': []}).status_code == 400
    assert system.post(handouts(space) + '/snapshots', json={'ids': ['1']}).status_code == 400
    document = sample_document(taken[1], [taken[0]])
    created = system.post(handouts(space), json={'title': '快照讲义', 'document': document})
    assert created.status_code == 200, created.text
    hid = created.json()['id']
    stored = json.loads(store.one('SELECT document FROM handouts WHERE id=?', (hid,))['document'])
    assert '![](/static/uploads/fig-a.png)' in stored['sections'][0]['blocks'][4]['question']['content']
    assert f'/static/uploads/{space}/fig-a.png' in created.json()['document']['sections'][0]['blocks'][4]['question']['content']
    assert system.get(f'{handouts(space)}/{hid}/changes').json() == {'changed': [], 'deleted': []}
    # A later bank edit does not touch the handout; it is reported for the teacher to decide.
    moved = system.post(f'/bank/{space}/api/questions/{first["id"]}/figure_align', data={'figure_align': 'center'})
    assert moved.status_code == 200, moved.text
    assert system.delete(f'/bank/{space}/api/questions/{second["id"]}').status_code == 200
    assert system.get(f'{handouts(space)}/{hid}/changes').json() == {'changed': [first['id']], 'deleted': [second['id']]}
    kept = system.get(f'{handouts(space)}/{hid}').json()['document']['sections'][0]['blocks'][4]['question']
    assert kept['figure_align'] == 'right' and kept['answer_markdown'] == '$x=1$'


def test_lesson_folders_hold_handouts(system):  # noqa: F811
    space = new_space(system, '讲义目录')['id']
    paper_folder = system.post(f'/platform/api/spaces/{space}/folders', json={'name': '试卷', 'kind': 'paper'}).json()['id']
    unit = system.post(f'/platform/api/spaces/{space}/folders', json={'name': '第一章', 'kind': 'lesson'}).json()['id']
    lesson = system.post(f'/platform/api/spaces/{space}/folders',
                         json={'name': '第一节', 'kind': 'lesson', 'parent_id': unit}).json()['id']
    assert system.post(handouts(space), json={'title': '错放', 'folder_id': paper_folder}).status_code == 400
    hid = system.post(handouts(space), json={'title': '第一节讲义', 'folder_id': lesson}).json()['id']
    assert system.get(f'{handouts(space)}/{hid}').json()['folder_id'] == lesson
    # Deleting a folder moves its handouts up a level, like papers.
    assert system.delete(f'/platform/api/spaces/{space}/folders/{lesson}').json()['papers_moved'] == 1
    assert system.get(f'{handouts(space)}/{hid}').json()['folder_id'] == unit
    assert system.patch(f'{handouts(space)}/{hid}', json={'folder_id': 0}).json()['folder_id'] is None
    assert system.patch(f'{handouts(space)}/{hid}', json={'folder_id': paper_folder}).status_code == 400
    assert system.patch(f'{handouts(space)}/{hid}', json={'folder_id': unit}).json()['folder_id'] == unit


def docx_text(blob):
    document = Document(BytesIO(blob))
    return ''.join(node.text or '' for node in document.element.body.iter(qn('w:t'))), document


def test_word_export_mirrors_the_template(system):  # noqa: F811
    space = new_space(system, '讲义导出')['id']
    picture(space, 'hd-fig.png')
    example = question(system, space, content='如图，求阴影面积（ ）。\n\n![](/static/uploads/hd-fig.png)\n'
                       '\\begin{choices}\\item $1$\\item $2$\\item $3$\\item $4$\\end{choices}',
                       question_type='single_choice', answer_markdown='选 B', figure_align='right',
                       image_paths=json.dumps(['/static/uploads/hd-fig.png'])).json()['question']
    exercise = question(system, space, content='计算 $2+3$。').json()['question']
    taken = snapshots(system, space, [example['id'], exercise['id']])
    hid = system.post(handouts(space), json={'title': '导出讲义', 'document': sample_document(taken[0], [taken[1]])}).json()['id']
    for version in ('student', 'teacher'):
        response = system.get(f'{handouts(space)}/{hid}/export', params={'format': 'docx', 'version': version})
        assert response.status_code == 200, response.text
        name = unquote(response.headers['content-disposition'].split("''", 1)[1])
        assert name == f"导出讲义（{'教师版' if version == 'teacher' else '学生版'}）.docx"
        text, document = docx_text(response.content)
        for expected in ('第 3 讲', '导出讲义', '学习目标', '定义 1.1  正弦', '定理 1.1', '易错', '例 1.1  基础',
                         '基础练习', '本节小结', '1  概念', '2  练习', '证明'):
            assert expected in text, expected
        header = ''.join(p.text for p in document.sections[0].header.paragraphs)
        assert '导出讲义' in header and '讲义' in header
        if version == 'teacher':
            assert '参考答案' in text and '由定义得' in text and '教师版' in text
        else:
            assert '参考答案' not in text and '由定义得' not in text and '学生版' in text
        # As in the PDF: the example's stem and its figure side by side, the choices below at full width.
        beside = [t for t in document.tables if len(t.columns) == 2 and t.cell(0, 0).text.startswith('例 1.1')]
        assert beside and beside[0].cell(0, 1)._tc.xpath('.//w:drawing')
        assert '<wp:anchor' not in zipfile.ZipFile(BytesIO(response.content)).read('word/document.xml').decode()
        assert any(t.cell(0, 0).text.strip().startswith('A.') and len(t.columns) == 4 for t in document.tables)
    assert system.get(f'{handouts(space)}/{hid}/export', params={'format': 'html'}).status_code == 400


def test_each_kind_counts_on_its_own_within_a_section(system):  # noqa: F811
    """定义 1.1, 例 1.1, 性质 1.1, 例 1.2, 定义 1.2, 图 1.1; a new section starts again at 2.1."""
    space = new_space(system, '讲义编号')['id']
    image = picture(space, 'hd-number.png')
    first, second = (question(system, space, content=f'求第 {n} 个值。').json()['question'] for n in (1, 2))
    one, two = snapshots(system, space, [first['id'], second['id']])
    example = {'kind': 'example', 'title': '', 'solution': '', 'space': 2}
    document = {'kicker': '', 'subtitle': '', 'objectives': [], 'settings': {'student_solutions': 'blank'}, 'sections': [
        {'title': '集合', 'blocks': [
            {'kind': 'definition', 'title': '集合', 'body': '对象的全体。'}, {**example, 'question': one},
            {'kind': 'property', 'title': '', 'body': '确定性。'}, {**example, 'question': two},
            {'kind': 'definition', 'title': '子集', 'body': '包含关系。'},
            {'kind': 'figure', 'image': image, 'caption': '韦恩图', 'size': 'small'}]},
        {'title': '运算', 'blocks': [{'kind': 'definition', 'title': '交集', 'body': '公共元素。'}]}]}
    hid = system.post(handouts(space), json={'title': '编号', 'document': document}).json()['id']
    response = system.get(f'{handouts(space)}/{hid}/export', params={'format': 'docx', 'version': 'student'})
    assert response.status_code == 200, response.text
    text, _ = docx_text(response.content)
    order = ['定义 1.1', '例 1.1', '性质 1.1', '例 1.2', '定义 1.2', '图 1.1', '定义 2.1']
    positions = [text.find(label) for label in order]
    assert -1 not in positions and positions == sorted(positions), dict(zip(order, positions))
    assert not any(old in text for old in ('性质 1.3', '例 1.4', '定义 1.5', '图 1.6'))  # the former shared count


def test_latex_source_uses_the_template_blocks(tmp_path):
    from platform_app import handout_export
    snapshot = {'source_id': 1, 'content': '下列正确的是（ ）。\\begin{choices}\\item 甲\\item 乙\\end{choices}',
                'question_type': 'single_choice', 'answer_markdown': '甲', 'image_layouts': {}, 'figure_align': 'right'}
    document = sample_document(snapshot, [dict(snapshot, source_id=2)])
    engine = SimpleNamespace(UPLOAD_DIR=str(tmp_path), UPLOAD_DIR_REL='static/uploads')

    def render(version):
        payload = {'title': '三角函数 & 应用', 'document': document, 'version': version,
                   'subject': '数学', 'stage': '高中', 'textbook': '人教A版'}
        return handout_export.Latex(engine, payload).render()

    student, teacher = render('student'), render('teacher')
    assert r'\usepackage{mathbank-handout}' in student and r'\documentclass[fontset=fandol]{exam-zh}' in student
    assert r'\begin{mbdefinition}{正弦}' in student and r'\textbf{正弦}' in student and r'\begin{itemize}' in student
    assert r'\begin{mbtheorem}{theorem}{定理}{}' in student and r'\begin{mbnote}{易错}{mbAccent}' in student
    assert r'100\% 的学生 \& 老师' in student and r'三角函数 \& 应用' in student
    assert r'\mbexamplelabel{基础}' in student and r'\mbanswerspace{3}' in student
    # \paren supplies the brackets; the stem's own empty pair is dropped.
    assert '（ ）' not in student and 'paren/show-answer=false' in student
    assert r'\begin{mbsolution}' not in student and r'\mbanswersheading' not in student
    assert r'\begin{mbsolution}' in teacher and r'\fillin{}' in teacher and r'\mbanswersheading' in teacher
    assert 'paren/show-answer=true' in teacher


def test_handouts_table_added_to_an_existing_install(tmp_path, monkeypatch):
    monkeypatch.setattr(store, 'DATA', tmp_path)
    store.initialize()
    with sqlite3.connect(tmp_path / 'platform.db') as c:
        c.execute('DROP TABLE handouts')
        c.execute("DELETE FROM platform_migrations WHERE name='handouts-v1'")
    store.initialize()
    assert store.rows('SELECT count(*) AS n FROM handouts') == [{'n': 0}]
    assert {r['name'] for r in store.rows('SELECT name FROM platform_migrations')} == set(migrations.NAMES)
    assert len(list((tmp_path / 'migration_backups').glob('before-add-handouts-*.db'))) == 1
    store.initialize()
    assert len(list((tmp_path / 'migration_backups').glob('before-add-handouts-*.db'))) == 1
